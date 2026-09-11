"""Быстрая запись клиента организацией + публичное подтверждение визита."""

from __future__ import annotations

import secrets
from datetime import timedelta

from django.conf import settings
from django.contrib.auth import get_user_model
from django.db import transaction
from django.utils import timezone
from django.utils.dateparse import parse_datetime
from rest_framework import permissions, status
from rest_framework.response import Response
from rest_framework.views import APIView

from catalog.models import Service

from .booking_actions import client_display_name, format_booking_when
from .booking_windows import book_time_window, resolve_selected_options
from .models import Booking, ProviderStaff
from .phone_clients import (
    client_brief,
    get_or_create_client_by_name,
    get_or_create_client_by_phone,
    normalize_phone,
    phone_digits,
)
from .serializers import BookingSerializer

User = get_user_model()


def _provider_context(user):
    """Вернуть (provider_id, is_allowed) для provider/staff с manage_bookings."""
    if not user or not getattr(user, "is_authenticated", False):
        return None, False
    if user.role == "provider":
        return user.id, True
    if user.role == "staff":
        link = (
            ProviderStaff.objects.filter(
                staff=user,
                is_active=True,
                invitation_status=ProviderStaff.InvitationStatus.ACCEPTED,
            )
            .order_by("id")
            .first()
        )
        if not link:
            return None, False
        perms = link.permissions or {}
        if not perms.get("manage_bookings", True):
            return None, False
        return link.provider_id, True
    return None, False


def ensure_client_confirm_token(booking: Booking) -> str:
    token = (booking.client_confirm_token or "").strip()
    if token:
        return token
    token = secrets.token_urlsafe(24)
    booking.client_confirm_token = token
    booking.save(update_fields=["client_confirm_token"])
    return token


def build_client_confirm_url(booking: Booking) -> str:
    front = (getattr(settings, "FRONTEND_URL", None) or "https://vsevmeste.space").rstrip("/")
    token = ensure_client_confirm_token(booking)
    return f"{front}/?visit_confirm={token}"


class ClientPhoneLookupView(APIView):
    """GET /api/booking/clients/lookup/?q=…|phone=…|name=… — поиск клиентов (подсказки)."""

    permission_classes = [permissions.IsAuthenticated]

    def get(self, request):
        from .phone_clients import search_clients_for_provider

        provider_id, ok = _provider_context(request.user)
        if not ok:
            return Response(status=status.HTTP_403_FORBIDDEN)
        q = (
            (request.query_params.get("q") or "").strip()
            or (request.query_params.get("phone") or "").strip()
            or (request.query_params.get("name") or "").strip()
        )
        if len(q) < 2:
            return Response({"results": [], "found": False, "client": None, "normalized_phone": ""})

        # Обратная совместимость: один клиент по точному телефону — только из своей базы
        phone_only = (request.query_params.get("phone") or "").strip()
        if phone_only and len(phone_digits(phone_only)) >= 10 and not (request.query_params.get("q") or "").strip():
            users = search_clients_for_provider(provider_id, phone_only, limit=1)
            if not users:
                return Response(
                    {
                        "found": False,
                        "client": None,
                        "results": [],
                        "normalized_phone": normalize_phone(phone_only),
                    }
                )
            brief = client_brief(users[0], request=request, provider_id=provider_id)
            return Response(
                {
                    "found": True,
                    "client": brief,
                    "results": [brief],
                    "normalized_phone": normalize_phone(phone_only),
                }
            )

        users = search_clients_for_provider(provider_id, q, limit=12)
        results = [client_brief(u, request=request, provider_id=provider_id) for u in users]
        return Response(
            {
                "found": bool(results),
                "client": results[0] if results else None,
                "results": results,
                "normalized_phone": normalize_phone(q) if phone_digits(q) else "",
            }
        )


class ProviderClientListView(APIView):
    """База клиентов: список, создание вручную, импорт Excel, удаление из базы."""

    permission_classes = [permissions.IsAuthenticated]

    def get(self, request):
        from .phone_clients import list_clients_for_provider

        provider_id, ok = _provider_context(request.user)
        if not ok:
            return Response({"detail": "Нет доступа"}, status=status.HTTP_403_FORBIDDEN)
        try:
            page = int(request.query_params.get("page") or 1)
        except (TypeError, ValueError):
            page = 1
        try:
            page_size = int(request.query_params.get("page_size") or 20)
        except (TypeError, ValueError):
            page_size = 20
        q = (request.query_params.get("q") or "").strip()
        data = list_clients_for_provider(
            provider_id, q=q, page=page, page_size=page_size, request=request
        )
        return Response(data)

    def post(self, request):
        from .models import ProviderClientCard
        from .phone_clients import client_brief, get_or_create_client_by_name, get_or_create_client_by_phone

        provider_id, ok = _provider_context(request.user)
        if not ok:
            return Response({"detail": "Нет доступа"}, status=status.HTTP_403_FORBIDDEN)

        # Импорт Excel/CSV
        upload = request.FILES.get("file") or request.FILES.get("excel")
        if upload:
            return self._import_file(request, provider_id, upload)

        data = request.data if hasattr(request.data, "get") else {}
        name = (data.get("name") or data.get("guest_name") or "").strip()
        phone = (data.get("phone") or "").strip()
        source = (data.get("acquisition_source") or data.get("source") or "").strip()[:120]
        if not name and not phone:
            return Response(
                {"detail": "Укажите имя или телефон."},
                status=status.HTTP_400_BAD_REQUEST,
            )
        try:
            if phone and len(phone_digits(phone)) >= 10:
                client = get_or_create_client_by_phone(phone=phone, name=name)
            else:
                client = get_or_create_client_by_name(name=name or "Клиент", phone=phone)
        except ValueError as e:
            return Response({"detail": str(e)}, status=status.HTTP_400_BAD_REQUEST)

        card, _ = ProviderClientCard.objects.get_or_create(provider_id=provider_id, client=client)
        if card.hidden:
            card.hidden = False
        if source:
            card.acquisition_source = source
        card.save()
        brief = client_brief(client, request=request, provider_id=provider_id)
        brief["has_memory"] = True
        brief["is_blocked"] = bool(card.is_blocked)
        brief["no_show_count"] = int(card.no_show_count or 0)
        brief["acquisition_source"] = card.acquisition_source or ""
        return Response(brief, status=status.HTTP_201_CREATED)

    def delete(self, request):
        from .models import ProviderClientCard

        provider_id, ok = _provider_context(request.user)
        if not ok:
            return Response({"detail": "Нет доступа"}, status=status.HTTP_403_FORBIDDEN)
        raw = request.query_params.get("client") or request.data.get("client")
        try:
            client_id = int(raw)
        except (TypeError, ValueError):
            return Response({"detail": "Укажите client."}, status=status.HTTP_400_BAD_REQUEST)
        card, _ = ProviderClientCard.objects.get_or_create(
            provider_id=provider_id, client_id=client_id
        )
        card.hidden = True
        card.save(update_fields=["hidden", "updated_at"])
        return Response({"ok": True, "hidden": True})

    def _import_file(self, request, provider_id, upload):
        ok, payload = import_clients_from_upload(provider_id, upload)
        if not ok:
            return Response(payload, status=status.HTTP_400_BAD_REQUEST)
        return Response(payload)


def import_clients_from_upload(provider_id, upload):
    """Импорт Excel/CSV в базу клиентов. Returns (ok: bool, payload: dict)."""
    from .models import ProviderClientCard
    from .phone_clients import get_or_create_client_by_name, get_or_create_client_by_phone

    name = (getattr(upload, "name", None) or "").lower()
    rows = []
    try:
        if name.endswith(".csv") or name.endswith(".txt"):
            import csv
            import io

            text = upload.read().decode("utf-8-sig", errors="replace")
            reader = csv.DictReader(io.StringIO(text))
            rows = list(reader)
        else:
            from openpyxl import load_workbook
            import io

            wb = load_workbook(io.BytesIO(upload.read()), read_only=True, data_only=True)
            ws = wb.active
            cells = list(ws.iter_rows(values_only=True))
            if not cells:
                return False, {"detail": "Пустой файл."}
            headers = [str(h or "").strip().lower() for h in cells[0]]
            for row in cells[1:]:
                item = {}
                for i, h in enumerate(headers):
                    if i < len(row):
                        item[h] = row[i]
                rows.append(item)
    except Exception as e:
        return False, {"detail": f"Не удалось прочитать файл: {e}"}

    created = 0
    updated = 0
    errors = []
    for idx, row in enumerate(rows, start=2):
        mapped = {}
        for k, v in (row or {}).items():
            key = str(k or "").strip().lower()
            mapped[key] = "" if v is None else str(v).strip()
        name_v = (
            mapped.get("name")
            or mapped.get("имя")
            or mapped.get("фио")
            or " ".join(
                p
                for p in (
                    mapped.get("фамилия") or mapped.get("last_name"),
                    mapped.get("имя") or mapped.get("first_name"),
                    mapped.get("отчество") or mapped.get("patronymic"),
                )
                if p
            ).strip()
        )
        if not name_v or name_v == (mapped.get("имя") or ""):
            fn = mapped.get("first_name") or mapped.get("имя") or ""
            ln = mapped.get("last_name") or mapped.get("фамилия") or ""
            pn = mapped.get("patronymic") or mapped.get("отчество") or ""
            name_v = " ".join(p for p in (ln, fn, pn) if p).strip() or name_v
        phone_v = mapped.get("phone") or mapped.get("телефон") or mapped.get("тел") or ""
        source_v = (
            mapped.get("source")
            or mapped.get("acquisition_source")
            or mapped.get("источник")
            or ""
        )[:120]
        if not name_v and not phone_v:
            continue
        try:
            if phone_v and len(phone_digits(phone_v)) >= 10:
                client = get_or_create_client_by_phone(phone=phone_v, name=name_v)
            else:
                client = get_or_create_client_by_name(name=name_v or "Клиент", phone=phone_v)
            card, was_created = ProviderClientCard.objects.get_or_create(
                provider_id=provider_id, client=client
            )
            changed = False
            if card.hidden:
                card.hidden = False
                changed = True
            if source_v and not card.acquisition_source:
                card.acquisition_source = source_v
                changed = True
            if changed:
                card.save()
            if was_created:
                created += 1
            else:
                updated += 1
        except Exception as e:
            errors.append({"row": idx, "detail": str(e)})

    return True, {
        "ok": True,
        "created": created,
        "updated": updated,
        "errors": errors[:20],
        "detail": f"Импортировано: новых {created}, обновлено {updated}.",
    }


class ClientMigrateRequestView(APIView):
    """GET/POST /api/booking/clients/migrate-request/ — заявка на перенос базы."""

    permission_classes = [permissions.IsAuthenticated]

    STATUS_LABELS = {
        "new": "Новая",
        "in_progress": "В работе",
        "done": "Готово",
        "rejected": "Отклонена",
    }

    def _serialize(self, obj):
        return {
            "id": obj.id,
            "status": obj.status,
            "status_label": self.STATUS_LABELS.get(obj.status, obj.status),
            "source_note": obj.source_note or "",
            "result_detail": obj.result_detail or "",
            "has_file": bool(obj.file),
            "file_name": (obj.file.name.rsplit("/", 1)[-1] if obj.file else "") or "",
            "created_at": obj.created_at.isoformat() if obj.created_at else None,
            "updated_at": obj.updated_at.isoformat() if obj.updated_at else None,
        }

    def get(self, request):
        from .models import ClientMigrateRequest

        provider_id, ok = _provider_context(request.user)
        if not ok:
            return Response({"detail": "Нет доступа"}, status=status.HTTP_403_FORBIDDEN)
        qs = ClientMigrateRequest.objects.filter(provider_id=provider_id).order_by("-created_at")[:20]
        results = [self._serialize(r) for r in qs]
        latest = results[0] if results else None
        return Response({"results": results, "latest": latest})

    def post(self, request):
        from .models import ClientMigrateRequest

        provider_id, ok = _provider_context(request.user)
        if not ok:
            return Response({"detail": "Нет доступа"}, status=status.HTTP_403_FORBIDDEN)

        upload = request.FILES.get("file") or request.FILES.get("excel")
        note = (
            (request.data.get("source_note") or request.data.get("note") or request.data.get("comment") or "")
            .strip()
        )[:2000]
        if not upload and not note:
            return Response(
                {"detail": "Прикрепите файл или опишите, откуда переносить базу."},
                status=status.HTTP_400_BAD_REQUEST,
            )

        obj = ClientMigrateRequest(provider_id=provider_id, source_note=note)
        if upload:
            obj.file = upload
        obj.save()
        return Response(self._serialize(obj), status=status.HTTP_201_CREATED)


def client_is_blocked_for_provider(provider_id: int, client_id: int) -> bool:
    from .phone_clients import client_is_blocked_for_provider as _fn

    return _fn(provider_id, client_id)


class BookForClientView(APIView):
    """POST /api/booking/book-for-client/ — мастер записывает клиента на свободное окно."""

    permission_classes = [permissions.IsAuthenticated]

    @transaction.atomic
    def post(self, request):
        provider_id, ok = _provider_context(request.user)
        if not ok:
            return Response(status=status.HTTP_403_FORBIDDEN)

        data = request.data or {}
        phone = (data.get("phone") or data.get("guest_phone") or "").strip()
        client_id = data.get("client") or data.get("client_id")
        name = (data.get("name") or data.get("guest_name") or "").strip()
        comment = (data.get("comment") or "")[:250]
        service_id = data.get("service") or data.get("service_id")
        starts_raw = data.get("starts_at")
        ends_raw = data.get("ends_at")
        staff_id = data.get("staff") or data.get("staff_id")
        option_ids = data.get("option_ids") or data.get("options") or []

        client = None
        if client_id not in (None, "", "null"):
            try:
                client = User.objects.filter(pk=int(client_id), role=User.Role.CLIENT).first()
            except (TypeError, ValueError):
                client = None
        if not client and phone and len(phone_digits(phone)) >= 10:
            try:
                client = get_or_create_client_by_phone(phone=phone, name=name)
            except ValueError as e:
                return Response({"detail": str(e)}, status=status.HTTP_400_BAD_REQUEST)
        if not client and name:
            try:
                client = get_or_create_client_by_name(name=name, phone=phone)
            except ValueError as e:
                return Response({"detail": str(e)}, status=status.HTTP_400_BAD_REQUEST)
        if not client:
            return Response(
                {"detail": "Выберите клиента из базы или укажите имя."},
                status=status.HTTP_400_BAD_REQUEST,
            )
        if client_is_blocked_for_provider(provider_id, client.id):
            return Response(
                {
                    "detail": "Клиент в чёрном списке. Снимите блокировку в базе клиентов или запишите только с предоплатой вне онлайн-записи.",
                },
                status=status.HTTP_400_BAD_REQUEST,
            )
        if name:
            from .phone_clients import _apply_name

            _apply_name(client, name)

        if not service_id or not starts_raw:
            return Response(
                {"detail": "Укажите service и starts_at."},
                status=status.HTTP_400_BAD_REQUEST,
            )

        starts_at = parse_datetime(str(starts_raw))
        if not starts_at:
            return Response({"detail": "Некорректное время начала."}, status=status.HTTP_400_BAD_REQUEST)
        if timezone.is_naive(starts_at):
            starts_at = timezone.make_aware(starts_at, timezone.get_current_timezone())

        try:
            service = Service.objects.get(pk=int(service_id), provider_id=provider_id, is_active=True)
        except (Service.DoesNotExist, TypeError, ValueError):
            return Response({"detail": "Услуга не найдена."}, status=status.HTTP_400_BAD_REQUEST)

        snapshots = resolve_selected_options(service, option_ids)
        extra = sum(int(o.get("extra_minutes") or 0) for o in snapshots)
        duration = int(service.duration_minutes or 30) + extra

        if ends_raw:
            ends_at = parse_datetime(str(ends_raw))
            if not ends_at:
                return Response({"detail": "Некорректное время окончания."}, status=status.HTTP_400_BAD_REQUEST)
            if timezone.is_naive(ends_at):
                ends_at = timezone.make_aware(ends_at, timezone.get_current_timezone())
        else:
            ends_at = starts_at + timedelta(minutes=duration)

        if ends_at <= starts_at:
            return Response({"detail": "Интервал слишком короткий."}, status=status.HTTP_400_BAD_REQUEST)

        sid = None
        if staff_id not in (None, "", "null", "any"):
            try:
                sid = int(staff_id)
            except (TypeError, ValueError):
                sid = None

        try:
            booking = book_time_window(
                provider_id,
                int(service_id),
                starts_at,
                ends_at,
                sid,
                client,
                comment,
                selected_options=snapshots,
                notify=True,
            )
        except ValueError as e:
            return Response({"detail": str(e)}, status=status.HTTP_400_BAD_REQUEST)

        booking = (
            Booking.objects.select_related("client", "provider", "service", "slot", "staff")
            .filter(pk=booking.pk)
            .first()
            or booking
        )
        ser = BookingSerializer(booking, context={"request": request})
        data_out = dict(ser.data)
        data_out["client_brief"] = client_brief(client)
        data_out["confirm_url"] = build_client_confirm_url(booking)
        return Response(data_out, status=status.HTTP_201_CREATED)


class PublicVisitConfirmView(APIView):
    """GET/POST /api/booking/public/visit-confirm/<token>/ — клиент подтверждает визит по ссылке."""

    permission_classes = [permissions.AllowAny]

    def _booking(self, token: str):
        token = (token or "").strip()
        if not token:
            return None
        return (
            Booking.objects.select_related("client", "provider", "service", "slot")
            .filter(client_confirm_token=token)
            .exclude(status=Booking.Status.CANCELLED)
            .first()
        )

    def get(self, request, token: str):
        booking = self._booking(token)
        if not booking:
            return Response({"detail": "Ссылка недействительна."}, status=status.HTTP_404_NOT_FOUND)
        org = (getattr(booking.provider, "organization_name", None) or "").strip() or "Организация"
        return Response(
            {
                "booking_id": booking.id,
                "status": booking.status,
                "already_confirmed": bool(booking.client_confirmed_at)
                or booking.status == Booking.Status.CONFIRMED,
                "org": org,
                "service": getattr(booking.service, "name", "") or "",
                "when": format_booking_when(booking),
                "client_name": client_display_name(booking.client),
            }
        )

    def post(self, request, token: str):
        booking = self._booking(token)
        if not booking:
            return Response({"detail": "Ссылка недействительна."}, status=status.HTTP_404_NOT_FOUND)
        if booking.status in (Booking.Status.DONE, Booking.Status.NO_SHOW, Booking.Status.CANCELLED):
            return Response({"detail": "Запись уже завершена."}, status=status.HTTP_400_BAD_REQUEST)
        if booking.payment_status == "pending":
            return Response(
                {"detail": "Сначала нужно внести предоплату."},
                status=status.HTTP_400_BAD_REQUEST,
            )

        now = timezone.now()
        update_fields = []
        if not booking.client_confirmed_at:
            booking.client_confirmed_at = now
            update_fields.append("client_confirmed_at")
        if booking.status == Booking.Status.NEW:
            booking.status = Booking.Status.CONFIRMED
            update_fields.append("status")
        if update_fields:
            booking.save(update_fields=update_fields)

        try:
            from notifications.delivery import deliver_booking_event

            when = format_booking_when(booking)
            name = client_display_name(booking.client)
            body = f"{name} подтвердил(а) визит · {when}" if name else f"Клиент подтвердил визит · {when}"
            deliver_booking_event(
                booking,
                "client_confirm",
                body,
                audience="org",
                title_org="Клиент подтвердил визит",
            )
        except Exception:
            pass

        return Response(
            {
                "ok": True,
                "status": booking.status,
                "already_confirmed": True,
                "org": (getattr(booking.provider, "organization_name", None) or "").strip() or "Организация",
                "service": getattr(booking.service, "name", "") or "",
                "when": format_booking_when(booking),
            }
        )
